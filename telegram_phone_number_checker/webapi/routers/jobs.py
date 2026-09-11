"""Job management + export endpoints for the web UI."""

import copy
import csv
import io
import os
import tempfile
from pathlib import Path
from typing import List, Optional

from fastapi import (APIRouter, Depends, File, Form, HTTPException, Request,
                     UploadFile)
from fastapi.responses import FileResponse
from pydantic import BaseModel
from starlette.background import BackgroundTask
from openpyxl import load_workbook

from ...exporters.csv_exporter import CsvExporter
from ...exporters.json_exporter import JsonExporter
from ...job_manager import JobController
from ...main import _create_job_with_numbers
from ...models import mask_phone, now_iso
from ...phone_utils import PhoneNormalizationError, normalize_phone
from ...rate_limiter import account_key_from_phone
from ...repositories.job_repository import JobRepository
from ...repositories.result_repository import ResultRepository
from ..deps import get_context, require_user
from ..runner import AutomationError

router = APIRouter(
    prefix="/api/jobs", tags=["jobs"], dependencies=[Depends(require_user)]
)

ITEM_COLUMNS = ResultRepository._ITEM_COLUMNS

try:
    MAX_UPLOAD_BYTES = max(1024, int(os.getenv("WEB_UI_MAX_UPLOAD_BYTES", str(5 * 1024 * 1024))))
except ValueError:
    MAX_UPLOAD_BYTES = 5 * 1024 * 1024
UPLOAD_CHUNK_BYTES = 64 * 1024
ALLOWED_UPLOAD_EXTENSIONS = {".csv", ".txt", ".xlsx"}
ALLOWED_UPLOAD_CONTENT_TYPES = {
    "text/csv", "text/plain", "application/csv",
    "application/vnd.ms-excel", "application/octet-stream",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class AccountBatchBody(BaseModel):
    telegram_account_id: str
    phone_numbers: str


class CreateJobBody(BaseModel):
    phone_numbers: str = ""
    job_name: Optional[str] = None
    max_attempts: Optional[int] = None
    telegram_account_id: Optional[str] = None
    mode: str = "SINGLE"
    account_batches: List[AccountBatchBody] = []
    auto_start: bool = False


def _job_detail(ctx, job_id: str) -> dict:
    job_repo = JobRepository(ctx.db)
    job = job_repo.get(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail=f"Không tìm thấy job {job_id}")
    from ...models import CheckStatus as CS

    result_repo = ResultRepository(ctx.db)
    if job.job_mode == "MULTI_PARENT":
        children = job_repo.list_children(job_id)
        branches = [_job_detail(ctx, child.id) for child in children]
        statuses = [b["status"] for b in branches]
        if any(s == "RUNNING" for s in statuses):
            parent_status = "RUNNING"
        elif any(s == "RATE_LIMITED" for s in statuses):
            parent_status = "RATE_LIMITED"
        elif statuses and all(s == "COMPLETED" for s in statuses):
            parent_status = "COMPLETED"
        elif statuses and all(s == "CANCELLED" for s in statuses):
            parent_status = "CANCELLED"
        elif any(s == "FAILED" for s in statuses):
            parent_status = "FAILED"
        elif any(s == "PAUSED" for s in statuses):
            parent_status = "PAUSED"
        else:
            parent_status = "CREATED"
        retry_times = [b.get("next_retry_at") for b in branches if b.get("next_retry_at")]
        return {
            "job_id": job.id, "name": job.name, "status": parent_status,
            "job_mode": "MULTI", "total": sum(b["total"] for b in branches),
            "processed": sum(b["processed"] for b in branches),
            "found": sum(b["found"] for b in branches),
            "not_discoverable": sum(b["not_discoverable"] for b in branches),
            "retry_queue": sum(b["retry_queue"] for b in branches),
            "errors": sum(b["errors"] for b in branches),
            "pending": sum(b["pending"] for b in branches),
            "active": ctx.runner.is_active(job_id) or any(b["active"] for b in branches),
            "created_at": job.created_at,
            "started_at": min((b["started_at"] for b in branches if b.get("started_at")), default=None),
            "finished_at": max((b["finished_at"] for b in branches if b.get("finished_at")), default=None) if branches and all(s in ("COMPLETED", "CANCELLED") for s in statuses) else None,
            "last_error_type": next((b["last_error_type"] for b in branches if b.get("last_error_type")), None),
            "last_error_message": next((b["last_error_message"] for b in branches if b.get("last_error_message")), None),
            "next_retry_at": min(retry_times) if retry_times else None,
            "account_blocked_until": None, "telegram_account_phone": None,
            "branches": branches,
        }
    account_blocked_until = None
    account_phone = job.telegram_account_phone or ctx.config.api_phone_number
    if account_phone:
        account_key = account_key_from_phone(account_phone)
        runtime = ctx.db.execute(
            "SELECT blocked_until FROM account_runtime_state WHERE account_key = ?",
            (account_key,),
        ).fetchone()
        if runtime:
            account_blocked_until = runtime["blocked_until"]
    return {
        "job_id": job.id,
        "name": job.name,
        "status": job.status.value,
        "total": job.total_items,
        "processed": job.processed_items,
        "found": result_repo.count_by_status(job_id, CS.FOUND),
        "not_discoverable": result_repo.count_by_status(job_id, CS.NOT_DISCOVERABLE),
        "retry_queue": result_repo.count_by_status(job_id, CS.RETRY_REQUIRED),
        "errors": result_repo.count_by_status(job_id, CS.PERMANENT_ERROR),
        "pending": result_repo.count_pending_all_statuses(job_id),
        "active": ctx.runner.is_active(job_id),
        "created_at": job.created_at,
        "started_at": job.started_at,
        "finished_at": job.finished_at,
        "last_error_type": job.last_error_type,
        "last_error_message": job.last_error_message,
        "next_retry_at": result_repo.get_next_retry_at(job_id),
        "account_blocked_until": account_blocked_until,
        "telegram_account_phone": mask_phone(job.telegram_account_phone) if job.telegram_account_phone else None,
        "job_mode": "SINGLE" if job.job_mode == "SINGLE" else "BRANCH",
        "branches": [],
    }


@router.get("")
async def list_jobs(request: Request, status: Optional[str] = None) -> dict:
    ctx = get_context(request)
    job_repo = JobRepository(ctx.db)
    jobs = job_repo.list_roots()
    items = []
    for job in jobs:
        detail = _job_detail(ctx, job.id)
        if status and detail["status"] != status.upper():
            continue
        items.append(detail)
    return {"items": items, "total": len(items)}


@router.post("")
async def create_job(request: Request, body: CreateJobBody) -> dict:
    ctx = get_context(request)
    mode = (body.mode or "SINGLE").upper()
    if mode not in ("SINGLE", "MULTI"):
        raise HTTPException(status_code=422, detail="mode phải là SINGLE hoặc MULTI.")
    if mode == "SINGLE":
        if not body.phone_numbers.strip():
            raise HTTPException(status_code=422, detail="Danh sách số điện thoại rỗng.")
        config = copy.copy(ctx.config)
        try:
            config.api_phone_number = ctx.account.phone_for(body.telegram_account_id)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc))
        if body.max_attempts is not None:
            config.max_attempts = body.max_attempts
        if body.auto_start and body.telegram_account_id:
            state = await ctx.account.status_for_id(body.telegram_account_id)
            if state.get("state") != "AUTHORIZED":
                raise HTTPException(status_code=409, detail="Tài khoản Telegram đã chọn chưa sẵn sàng để chạy job.")
        job_id = _create_job_with_numbers(ctx.db, config, body.phone_numbers, body.job_name)
    else:
        if len(body.account_batches) < 2:
            raise HTTPException(status_code=422, detail="Chế độ nhiều tài khoản cần ít nhất 2 tài khoản.")
        ids = [b.telegram_account_id for b in body.account_batches]
        if len(ids) != len(set(ids)):
            raise HTTPException(status_code=422, detail="Mỗi tài khoản chỉ được xuất hiện một lần trong job.")
        if any(not b.phone_numbers.strip() for b in body.account_batches):
            raise HTTPException(status_code=422, detail="Mỗi tài khoản phải có danh sách số riêng.")
        if body.auto_start:
            for batch in body.account_batches:
                state = await ctx.account.status_for_id(batch.telegram_account_id)
                if state.get("state") != "AUTHORIZED":
                    raise HTTPException(status_code=409, detail="Tất cả tài khoản Telegram phải ở trạng thái đã kết nối trước khi chạy song song.")
        repo = JobRepository(ctx.db)
        job_id = __import__("uuid").uuid4().hex[:12]
        parent_name = body.job_name or f"multi_{now_iso()}"
        total = 0
        # Parent + every branch + every item must either all exist or none exist.
        with ctx.db.transaction():
            repo.create(job_id, name=parent_name, total_items=0, job_mode="MULTI_PARENT")
            for index, batch in enumerate(body.account_batches, start=1):
                config = copy.copy(ctx.config)
                config.api_phone_number = ctx.account.phone_for(batch.telegram_account_id)
                if not config.api_phone_number:
                    raise ValueError("Không tìm thấy tài khoản Telegram đã chọn.")
                if body.max_attempts is not None:
                    config.max_attempts = body.max_attempts
                account = ctx.account.get_account(batch.telegram_account_id)
                label = account["label"] if account else f"Tài khoản {index}"
                child_id = _create_job_with_numbers(
                    ctx.db, config, batch.phone_numbers, f"{parent_name} · {label}",
                    job_mode="MULTI_BRANCH", parent_job_id=job_id,
                )
                child = repo.get(child_id)
                total += child.total_items if child else 0
            repo.update_totals(job_id, total, 0, 0, 0, 0, 0)
    if body.auto_start:
        try:
            ctx.runner.start(job_id)
        except (AutomationError, ValueError) as exc:
            raise HTTPException(status_code=409, detail=str(exc))
    return _job_detail(ctx, job_id)


def _phones_from_text(text: str, is_csv: bool) -> List[str]:
    if not is_csv:
        return [
            line.strip()
            for line in text.splitlines()
            if line.strip() and line.strip().lower() not in ("phone", "number")
        ]

    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return []
    header = [cell.strip().lower() for cell in rows[0]]
    phone_index = next(
        (i for i, name in enumerate(header) if name in ("phone", "number", "phone_number")),
        None,
    )
    data_rows = rows[1:] if phone_index is not None else rows
    index = phone_index if phone_index is not None else 0
    phones: List[str] = []
    for row in data_rows:
        if index >= len(row):
            continue
        value = row[index].strip()
        if value and value.lower() not in ("phone", "number", "phone_number"):
            phones.append(value)
    return phones


def _phones_from_xlsx(raw: bytes) -> List[str]:
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="File Excel không hợp lệ.") from exc
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        first = next(rows, None)
        if first is None:
            return []
        cells = ["" if v is None else str(v).strip() for v in first]
        header = [v.lower() for v in cells]
        phone_index = next((i for i, name in enumerate(header) if name in ("phone", "number", "phone_number", "số điện thoại")), None)
        index = phone_index if phone_index is not None else 0
        source_rows = rows if phone_index is not None else iter([first, *rows])
        phones: List[str] = []
        for row in source_rows:
            if index >= len(row) or row[index] is None:
                continue
            value = str(row[index]).strip()
            if value and value.lower() not in ("phone", "number", "phone_number", "số điện thoại"):
                phones.append(value)
        return phones
    finally:
        workbook.close()


def _unique_distribution_inputs(phones: List[str], region: str) -> List[str]:
    seen = set()
    result = []
    for raw in phones:
        value = str(raw).strip()
        if not value:
            continue
        try:
            normalized = normalize_phone(value, region)
            key = ("phone", normalized)
            output = normalized
        except PhoneNormalizationError:
            key = ("invalid", value.casefold())
            output = value
        if key in seen:
            continue
        seen.add(key)
        result.append(output)
    return result


def _create_distributed_import_job(ctx, account_ids: List[str], phones: List[str], job_name: Optional[str]) -> str:
    ids = list(dict.fromkeys(account_ids))
    if len(ids) < 2:
        raise HTTPException(status_code=422, detail="Cần chọn ít nhất 2 tài khoản để chia dữ liệu.")
    unique_inputs = _unique_distribution_inputs(phones, ctx.config.default_phone_region)
    if len(unique_inputs) < 2:
        raise HTTPException(status_code=422, detail="Cần ít nhất 2 target duy nhất để chia dữ liệu.")
    active_ids = ids[: min(len(ids), len(unique_inputs))]
    buckets = {account_id: [] for account_id in active_ids}
    for index, value in enumerate(unique_inputs):
        buckets[active_ids[index % len(active_ids)]].append(value)

    repo = JobRepository(ctx.db)
    parent_id = __import__("uuid").uuid4().hex[:12]
    parent_name = job_name or f"distributed_{now_iso()}"
    total = 0
    with ctx.db.transaction():
        repo.create(parent_id, name=parent_name, total_items=0, job_mode="MULTI_PARENT")
        for index, account_id in enumerate(active_ids, start=1):
            account = ctx.account.get_account(account_id)
            if account is None:
                raise HTTPException(status_code=400, detail=f"Không tìm thấy tài khoản #{index}.")
            config = copy.copy(ctx.config)
            config.api_phone_number = ctx.account.phone_for(account_id)
            child_id = _create_job_with_numbers(
                ctx.db, config, "\n".join(buckets[account_id]),
                f"{parent_name} · {account['label']}",
                job_mode="MULTI_BRANCH", parent_job_id=parent_id,
            )
            child = repo.get(child_id)
            total += child.total_items if child else 0
        repo.update_totals(parent_id, total, 0, 0, 0, 0, 0)
    return parent_id


@router.post("/import")
async def import_job(
    request: Request,
    file: UploadFile = File(...),
    job_name: Optional[str] = Form(None),
    telegram_account_id: Optional[str] = Form(None),
    telegram_account_ids: Optional[str] = Form(None),
) -> dict:
    ctx = get_context(request)
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in ALLOWED_UPLOAD_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Chỉ hỗ trợ file .csv, .txt hoặc .xlsx.")
    content_type = (file.content_type or "").split(";", 1)[0].strip().lower()
    if content_type and content_type not in ALLOWED_UPLOAD_CONTENT_TYPES:
        raise HTTPException(status_code=415, detail="Định dạng nội dung file không được hỗ trợ.")
    if file.size is not None and file.size > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File vượt quá giới hạn dung lượng upload.")

    chunks = []
    total = 0
    while True:
        chunk = await file.read(UPLOAD_CHUNK_BYTES)
        if not chunk:
            break
        total += len(chunk)
        if total > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="File vượt quá giới hạn dung lượng upload.")
        chunks.append(chunk)
    raw = b"".join(chunks)
    if suffix == ".xlsx":
        phones = _phones_from_xlsx(raw)
    else:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise HTTPException(status_code=400, detail="File CSV/TXT phải là UTF-8.")
        phones = _phones_from_text(text, suffix == ".csv")
    if not phones:
        raise HTTPException(
            status_code=422, detail="Không tìm thấy số điện thoại trong file."
        )

    multi_ids = [x.strip() for x in (telegram_account_ids or "").split(",") if x.strip()]
    if len(multi_ids) >= 2:
        job_id = _create_distributed_import_job(ctx, multi_ids, phones, job_name)
        return _job_detail(ctx, job_id)

    config = copy.copy(ctx.config)
    selected_id = multi_ids[0] if len(multi_ids) == 1 else telegram_account_id
    try:
        config.api_phone_number = ctx.account.phone_for(selected_id)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    job_id = _create_job_with_numbers(ctx.db, config, ", ".join(phones), job_name)
    return _job_detail(ctx, job_id)


@router.get("/{job_id}")
async def get_job(request: Request, job_id: str) -> dict:
    ctx = get_context(request)
    return _job_detail(ctx, job_id)


@router.get("/{job_id}/items")
async def get_items(
    request: Request, job_id: str, page: int = 1, page_size: int = 50,
    status: Optional[str] = None, q: Optional[str] = None,
) -> dict:
    ctx = get_context(request)
    repo = JobRepository(ctx.db)
    job = repo.get(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy job")
    page = max(1, page)
    page_size = min(max(1, page_size), 200)
    targets = repo.list_children(job_id) if job.job_mode == "MULTI_PARENT" else [job]
    target_ids = [target.id for target in targets]
    if not target_ids:
        return {"items": [], "total": 0, "page": page, "page_size": page_size}
    placeholders = ",".join("?" for _ in target_ids)
    where = [f"c.job_id IN ({placeholders})"]
    params: list = list(target_ids)
    if status:
        where.append("c.status = ?")
        params.append(status.upper())
    if q:
        where.append("(c.original_phone LIKE ? OR c.normalized_phone LIKE ?)")
        pattern = f"%{q}%"
        params.extend((pattern, pattern))
    sql_where = " AND ".join(where)
    total_row = ctx.db.execute(
        f"SELECT COUNT(*) AS c FROM check_items c WHERE {sql_where}", tuple(params)
    ).fetchone()
    total = total_row["c"] if total_row else 0
    offset = (page - 1) * page_size
    columns = ", ".join(f"c.{col.strip()}" for col in ITEM_COLUMNS.split(","))
    rows = ctx.db.execute(
        f"SELECT {columns}, j.telegram_account_phone AS branch_account_phone "
        f"FROM check_items c JOIN jobs j ON j.id = c.job_id WHERE {sql_where} "
        f"ORDER BY c.id LIMIT ? OFFSET ?", tuple(params + [page_size, offset]),
    ).fetchall()
    items = []
    for row in rows:
        d = dict(row)
        d["masked_phone"] = _mask(d.get("original_phone") or "")
        branch_phone = d.pop("branch_account_phone", None)
        d["telegram_account_phone"] = _mask(branch_phone) if branch_phone else None
        d.pop("original_phone", None)
        d.pop("normalized_phone", None)
        items.append(d)
    return {"items": items, "total": total, "page": page, "page_size": page_size}


@router.post("/{job_id}/start")
async def start_job(request: Request, job_id: str) -> dict:
    ctx = get_context(request)
    _job_detail(ctx, job_id)  # raises 404 when missing
    try:
        ctx.runner.start(job_id)
    except (AutomationError, ValueError) as exc:
        raise HTTPException(status_code=409, detail=str(exc))
    return {"ok": True, "job_id": job_id, "active": True}


@router.post("/{job_id}/pause")
async def pause_job(request: Request, job_id: str) -> dict:
    ctx = get_context(request)
    _job_detail(ctx, job_id)
    ctx.runner.pause(job_id)
    return {"ok": True, "job_id": job_id}


@router.post("/{job_id}/resume")
async def resume_job(request: Request, job_id: str) -> dict:
    ctx = get_context(request)
    _job_detail(ctx, job_id)
    try:
        ctx.runner.resume(job_id)
    except Exception as exc:
        raise HTTPException(status_code=409, detail=str(exc))
    return {"ok": True, "job_id": job_id, "active": True}


@router.post("/{job_id}/cancel")
async def cancel_job(request: Request, job_id: str) -> dict:
    ctx = get_context(request)
    _job_detail(ctx, job_id)
    await ctx.runner.cancel(job_id)
    return {"ok": True, "job_id": job_id}


@router.get("/{job_id}/export")
async def export_job(
    request: Request,
    job_id: str,
    format: str = "json",
) -> FileResponse:
    ctx = get_context(request)
    result_repo = ResultRepository(ctx.db)
    job_repo = JobRepository(ctx.db)
    job = job_repo.get(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Không tìm thấy job.")
    targets = job_repo.list_children(job_id) if job.job_mode == "MULTI_PARENT" else [job]
    rows = []
    for target in targets:
        for row in result_repo.all_items(target.id):
            row["telegram_account_phone"] = target.telegram_account_phone
            rows.append(row)
    if not rows:
        raise HTTPException(status_code=404, detail="Không có dữ liệu để export.")
    ext = ".csv" if format == "csv" else ".json"
    suffix = now_iso().replace(":", "-")[:19]
    out = Path(tempfile.mkdtemp()) / f"job_{job_id}_results{suffix}{ext}"
    if format == "csv":
        CsvExporter().export(rows, str(out))
        media = "text/csv; charset=utf-8"
    else:
        JsonExporter().export(rows, str(out))
        media = "application/json; charset=utf-8"
    return FileResponse(
        str(out),
        media_type=media,
        filename=out.name,
        background=BackgroundTask(os.unlink, str(out)),
    )


@router.delete("/{job_id}")
async def delete_job(request: Request, job_id: str) -> dict:
    ctx = get_context(request)
    _job_detail(ctx, job_id)
    repo = JobRepository(ctx.db)
    job = repo.get(job_id)
    targets = repo.list_children(job_id) if job and job.job_mode == "MULTI_PARENT" else ([job] if job else [])
    ids = [target.id for target in targets] + ([job_id] if job and job.job_mode == "MULTI_PARENT" else [])
    for target_id in ids:
        if ctx.runner.is_active(target_id) or repo.has_live_worker_lease(target_id):
            raise HTTPException(status_code=409, detail="Job hoặc một nhánh tài khoản đang chạy, không thể xóa.")
    with ctx.db.transaction():
        for target in targets:
            ctx.db.execute("DELETE FROM check_items WHERE job_id = ?", (target.id,))
        for target in targets:
            ctx.db.execute("DELETE FROM jobs WHERE id = ?", (target.id,))
        if job and job.job_mode == "MULTI_PARENT":
            ctx.db.execute("DELETE FROM jobs WHERE id = ?", (job_id,))
        elif job:
            ctx.db.execute("DELETE FROM check_items WHERE job_id = ?", (job_id,))
            ctx.db.execute("DELETE FROM jobs WHERE id = ?", (job_id,))
    return {"ok": True, "job_id": job_id}


def _mask(phone: str) -> str:
    from ...models import mask_phone

    return mask_phone(phone)
