import asyncio
import logging
import os
import uuid
from pathlib import Path
from typing import Optional

import click
from openpyxl import load_workbook

from .config import Config
from .database import Database
from .exporters.csv_exporter import CsvExporter
from .exporters.json_exporter import JsonExporter
from .job_manager import JobController, JobManager, JobPausedError
from .lease_keeper import LostOwnershipError
from .logging_config import log_event, setup_logging
from .models import JobStatus, now_iso
from .phone_utils import PhoneNormalizationError, normalize_phone
from .repositories.job_repository import JobRepository
from .repositories.result_repository import ResultRepository

# Backward-compatible re-exports (existing tests import from main)
from .telegram_service import get_human_readable_user_status, parse_proxy

logger = logging.getLogger(__name__)


def _make_db(config: Config) -> Database:
    config.ensure_dirs()
    return Database(config.db_path, database_url=config.database_url)


def _job_manager(config: Config) -> Database:
    return _make_db(config)


@click.group()
def cli():
    """Check whether phone numbers are connected to Telegram accounts."""
    setup_logging()


@cli.group()
def job():
    """Manage jobs: status, pause, resume."""


@job.command("status")
@click.argument("job_id")
@click.option("--db", "db_path", help="Database path", default=None)
def job_status(job_id: str, db_path: Optional[str]):
    """Show progress for a job."""
    config = Config.from_cli(database_path=db_path)
    db = _make_db(config)
    controller = JobController(db)
    info = controller.status(job_id)
    click.echo(f"Job: {info['job_id']}")
    click.echo(f"Name: {info['name'] or '-'}")
    click.echo(f"Status: {info['status']}")
    click.echo(f"Total: {info['total']:,}")
    click.echo(f"Processed: {info['processed']:,}")
    click.echo(f"Found: {info['found']:,}")
    click.echo(f"Not discoverable: {info['not_discoverable']:,}")
    click.echo(f"Retry queue: {info['retry_queue']:,}")
    click.echo(f"Errors: {info['errors']:,}")
    click.echo(f"Pending: {info['pending']:,}")
    db.close()


@job.command("pause")
@click.argument("job_id")
@click.option("--db", "db_path", help="Database path", default=None)
def job_pause(job_id: str, db_path: Optional[str]):
    """Pause a job (safe for the in-flight item)."""
    config = Config.from_cli(database_path=db_path)
    db = _make_db(config)
    JobController(db).pause(job_id)
    click.echo(f"Job {job_id} paused.")
    db.close()


@job.command("resume")
@click.argument("job_id")
@click.option("--db", "db_path", help="Database path", default=None)
def job_resume(job_id: str, db_path: Optional[str]):
    """Resume a paused job from its database checkpoint."""
    config = Config.from_cli(database_path=db_path)
    db = _make_db(config)
    JobController(db).resume(job_id)
    click.echo(f"Job {job_id} resumed from checkpoint.")
    db.close()


@job.command("validate")
@click.option("--db", "db_path", help="Database path", default=None)
def job_validate(db_path: Optional[str]):
    """Check the whole database for lifecycle/orphan-lease inconsistencies."""
    config = Config.from_cli(database_path=db_path)
    db = _make_db(config)
    from .database_validation import validate_database_state

    findings = validate_database_state(db)
    if not findings:
        click.echo("OK: database is consistent.")
    else:
        for severity, msg in findings:
            click.echo(f"[{severity}] {msg}")
        click.echo(f"{len(findings)} finding(s).")
        bad = [f for f in findings if f[0] == "ERROR"]
        if bad:
            raise click.ClickException(f"{len(bad)} invariant error(s) found")
    db.close()


@cli.command("check")
@click.option(
    "-p", "--phone-numbers", help="Comma-separated phone numbers to check", default=None
)
@click.option(
    "--job-id", "job_id", help="Reuse an existing job by ID to resume", default=None
)
@click.option("--job-name", "job_name", help="Name for the job", default=None)
@click.option("--db", "db_path", help="Database path", default=None)
@click.option(
    "--auto-resume",
    "auto_resume",
    is_flag=True,
    default=None,
    help="Auto resume a paused job",
)
@click.option("--api-id", help="Telegram API ID", envvar="API_ID")
@click.option("--api-hash", help="Telegram API hash", envvar="API_HASH")
@click.option(
    "--api-phone-number", help="Your Telegram phone number", envvar="PHONE_NUMBER"
)
@click.option("--proxy", help="Proxy URL", default=None, envvar="PROXY")
@click.option(
    "--min-request-interval",
    "min_interval",
    type=float,
    default=None,
    envvar="MIN_REQUEST_INTERVAL_SECONDS",
    help="Safety pacing between requests (seconds)",
)
@click.option(
    "--max-attempts",
    type=int,
    default=None,
    envvar="MAX_ATTEMPTS",
    help="Max retry attempts",
)
@click.option("--output", help="Export results to JSON after run", default=None)
def check_cmd(
    phone_numbers: Optional[str],
    job_id: Optional[str],
    job_name: Optional[str],
    db_path: Optional[str],
    auto_resume: Optional[bool],
    api_id: Optional[str],
    api_hash: Optional[str],
    api_phone_number: Optional[str],
    proxy: Optional[str],
    min_interval: Optional[float],
    max_attempts: Optional[int],
    output: Optional[str],
):
    """Check phone numbers against Telegram and store progress in SQLite."""
    config = Config.from_cli(
        api_id=api_id,
        api_hash=api_hash,
        api_phone_number=api_phone_number,
        proxy=proxy,
        database_path=db_path,
    )
    if min_interval is not None:
        config.min_request_interval_seconds = min_interval
    if max_attempts is not None:
        config.max_attempts = max_attempts

    db = _make_db(config)

    if job_id is None:
        if not phone_numbers:
            raise click.ClickException("Provide --phone-numbers or --job-id")
        job_id = _create_job_with_numbers(db, config, phone_numbers, job_name)
        click.echo(f"Created job {job_id}")

    asyncio.run(_run_job(config, db, job_id, auto_resume))

    if output:
        _export_job(db, job_id, output)
    db.close()


@cli.command("import")
@click.argument("source")
@click.option("--job-name", "job_name", help="Name for the new job", default=None)
@click.option("--db", "db_path", help="Database path", default=None)
@click.option("--max-attempts", type=int, default=None, envvar="MAX_ATTEMPTS")
def import_cmd(
    source: str,
    job_name: Optional[str],
    db_path: Optional[str],
    max_attempts: Optional[int],
):
    """Import phones from a CSV, TXT or XLSX file into a new job queue."""
    config = Config.from_cli(database_path=db_path)
    if max_attempts is not None:
        config.max_attempts = max_attempts
    db = _make_db(config)
    phones = _read_phone_file(source)
    job_id = _create_job_with_numbers(db, config, ", ".join(phones), job_name)
    total = JobRepository(db).get(job_id).total_items
    click.echo(f"Imported {total} unique phone numbers into job {job_id}")
    db.close()


@cli.command("web")
@click.option(
    "--host",
    default=None,
    envvar="WEB_UI_HOST",
    help="Host to bind (default 127.0.0.1)",
)
@click.option(
    "--port",
    default=None,
    type=int,
    envvar="WEB_UI_PORT",
    help="Port to bind (default 8000)",
)
@click.option(
    "--reload", is_flag=True, default=False, help="Enable dev auto-reload (uvicorn)"
)
def web_cmd(host: Optional[str], port: Optional[int], reload: bool):
    """Run the web UI (FastAPI API + served React build)."""
    import uvicorn

    bind_host = host or os.getenv("WEB_UI_HOST", "127.0.0.1")
    bind_port = port or int(os.getenv("WEB_UI_PORT", "8000"))
    if reload:
        uvicorn.run(
            "telegram_phone_number_checker.webapi.app:web",
            host=bind_host,
            port=bind_port,
            reload=reload,
        )
    else:
        from .webapi.app import web

        uvicorn.run(web, host=bind_host, port=bind_port)


@cli.command("export")
@click.argument("job_id")
@click.argument("output")
@click.option("--format", "fmt", type=click.Choice(["json", "csv"]), default="json")
@click.option("--db", "db_path", help="Database path", default=None)
def export_cmd(job_id: str, output: str, fmt: str, db_path: Optional[str]):
    """Export an existing job's results."""
    config = Config.from_cli(database_path=db_path)
    db = _make_db(config)
    _export_job(db, job_id, output, fmt=fmt)
    db.close()


def _read_phone_file(source: str) -> list:
    path = Path(source)
    if not path.exists():
        raise click.ClickException(f"File not found: {source}")
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            first = next(rows, None)
            if first is None:
                return []
            header = ["" if value is None else str(value).strip().lower() for value in first]
            index = next((i for i, name in enumerate(header) if name in ("phone", "number", "phone_number", "số điện thoại")), 0)
            data_rows = rows if index != 0 or (header and header[0] in ("phone", "number", "phone_number", "số điện thoại")) else iter([first, *rows])
            return [str(row[index]).strip() for row in data_rows if index < len(row) and row[index] is not None and _valid_phone_or_skip(str(row[index]).strip())]
        finally:
            workbook.close()
    if suffix not in {".csv", ".txt"}:
        raise click.ClickException("Only .csv, .txt and .xlsx files are supported")
    phones = []
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line:
            continue
        if suffix == ".csv" and "," in line:
            phones.extend(p for p in (part.strip() for part in line.split(",")) if _valid_phone_or_skip(p))
        elif _valid_phone_or_skip(line):
            phones.append(line)
    return phones


def _valid_phone_or_skip(phone: str) -> bool:
    if phone.lower() in ("phone", "number"):
        return False
    return bool(phone)


def _normalize_or_skip(raw: str, default_region: str = "VN") -> Optional[str]:
    try:
        return normalize_phone(raw, default_region=default_region)
    except PhoneNormalizationError:
        return None


def _create_job_with_numbers(
    db: Database, config: Config, phones_csv: str, job_name: Optional[str],
    *, job_mode: str = "SINGLE", parent_job_id: Optional[str] = None,
) -> str:
    job_repo = JobRepository(db)
    result_repo = ResultRepository(db)
    raw_phones = [
        token.strip()
        for line in phones_csv.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        for token in line.split(",")
        if token.strip()
    ]
    normalized = []
    invalid = []
    seen = set()
    for raw in raw_phones:
        norm = _normalize_or_skip(raw, config.default_phone_region)
        if norm is None:
            invalid.append(raw)
            continue
        if norm in seen:
            continue
        seen.add(norm)
        normalized.append((raw, norm))

    job_id = uuid.uuid4().hex[:12]
    if not job_name:
        job_name = f"job_{now_iso()}"
    # Job metadata and every imported item are one atomic unit. Repository
    # commit() calls are deferred by Database.transaction().
    with db.transaction():
        job_repo.create(
            job_id, name=job_name, total_items=len(normalized) + len(invalid),
            telegram_account_phone=config.api_phone_number, job_mode=job_mode,
            parent_job_id=parent_job_id,
        )
        if invalid:
            log_event(logger, "IMPORT_INVALID_PHONES", count=len(invalid))
        for raw, norm in normalized:
            result_repo.insert(job_id, raw, norm, config.max_attempts)
        for raw in invalid:
            result_repo.insert_invalid(job_id, raw)
        job_repo.update_totals(
            job_id, len(normalized) + len(invalid), len(invalid), 0, 0, len(invalid), 0
        )
    return job_id


async def _run_job(
    config: Config, db: Database, job_id: str, auto_resume: Optional[bool]
):
    manager = JobManager(config, db)
    job_repo = JobRepository(db)
    if auto_resume is None:
        auto_resume = config.auto_resume
    try:
        status = await manager.run(job_id, auto_resume=auto_resume)
        click.echo(f"Job {job_id} finished: {status.value}")
    except JobPausedError as e:
        click.echo(str(e))
    except asyncio.CancelledError:
        click.echo(f"Job {job_id} paused on shutdown.")
    except LostOwnershipError as e:
        # A stale worker must never overwrite the state of a successor that may
        # already own this job. JobManager deliberately leaves successor state
        # untouched when ownership is lost, and the CLI wrapper must do the same.
        log_event(logger, "JOB_LOST_OWNERSHIP", job_id=job_id)
        click.echo(f"Job {job_id} lost worker ownership: {e}", err=True)
    except Exception as e:
        # Fatal connection/auth/runtime failure. Only mark FAILED while the job
        # is currently unowned; this prevents a late CLI exception from
        # clobbering a newly claimed successor worker.
        log_event(logger, "JOB_FAILED", job_id=job_id, error_type=type(e).__name__)
        now = now_iso()
        db.execute(
            "UPDATE jobs SET status = ?, updated_at = ? "
            "WHERE id = ? AND worker_id IS NULL",
            (JobStatus.FAILED.value, now, job_id),
        )
        db.commit()
        click.echo(f"Job {job_id} failed: {e}", err=True)


def _export_job(db: Database, job_id: str, output: str, fmt: str = "json"):
    repo = ResultRepository(db)
    rows = repo.all_items(job_id)
    if not rows:
        raise click.ClickException(f"No items found for job {job_id}")
    if fmt == "csv":
        CsvExporter().export(rows, output)
    else:
        JsonExporter().export(rows, output)
    click.echo(f"Exported {len(rows)} rows to {output}")


if __name__ == "__main__":
    cli()


def main_entrypoint() -> None:
    """Console-script entry point referenced by pyproject.toml."""
    cli()
