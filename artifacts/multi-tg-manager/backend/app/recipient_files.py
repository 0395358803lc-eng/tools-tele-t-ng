"""Parse recipient lists from CSV/XLS/XLSX without persisting uploaded files."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

import xlrd
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls"}
USERNAME_RE = re.compile(r"^[A-Za-z0-9_]{5,32}$")
PHONE_CHARS_RE = re.compile(r"^[+\d\s().-]+$")

PHONE_HEADERS = {
    "phone", "phonenumber", "mobile", "tel", "telephone", "sodienthoai", "sdt",
    "dien thoai", "dienthoai", "recipientphone", "telegramphone",
}
USERNAME_HEADERS = {
    "username", "user", "usname", "telegramusername", "telegramuser", "handle",
    "accountusername", "recipientusername",
}
TARGET_HEADERS = {"target", "recipient", "receiver", "nguoinhan", "dichden"}


@dataclass(frozen=True)
class Recipient:
    value: str
    kind: str  # phone | username


def _ascii_key(value: object) -> str:
    raw = str(value or "").replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", raw).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _header_kind(value: object) -> str | None:
    key = _ascii_key(value)
    if key in {_ascii_key(v) for v in PHONE_HEADERS}:
        return "phone"
    if key in {_ascii_key(v) for v in USERNAME_HEADERS}:
        return "username"
    if key in {_ascii_key(v) for v in TARGET_HEADERS}:
        return "target"
    return None


def _normalize_tme(value: str) -> str | None:
    raw = value.strip()
    body = raw
    if not re.match(r"^[a-z]+://", body, flags=re.I):
        if body.lower().startswith(("t.me/", "telegram.me/")):
            body = "https://" + body
        else:
            return None
    try:
        parsed = urlparse(body)
    except Exception:
        return None
    host = parsed.netloc.lower()
    if host not in {"t.me", "www.t.me", "telegram.me", "www.telegram.me"}:
        return None
    seg = parsed.path.strip("/").split("/", 1)[0]
    if not seg or seg.startswith("+") or seg.lower() in {"c", "joinchat"}:
        return None
    return f"@{seg}" if USERNAME_RE.fullmatch(seg) else None


def normalize_recipient(value: object, hint: str | None = None, allow_bare_username: bool = False) -> Recipient | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = str(value).strip()
    if not text:
        return None

    tme = _normalize_tme(text)
    if tme:
        return Recipient(tme, "username")

    if text.startswith("@"):
        username = text[1:].strip()
        if USERNAME_RE.fullmatch(username):
            return Recipient(f"@{username}", "username")
        return None

    if hint in {"phone", "target"} or PHONE_CHARS_RE.fullmatch(text):
        digits = re.sub(r"\D", "", text)
        if 7 <= len(digits) <= 15:
            return Recipient(f"+{digits}", "phone")

    if hint in {"username", "target"} or allow_bare_username:
        username = text.strip().lstrip("@")
        if USERNAME_RE.fullmatch(username):
            return Recipient(f"@{username}", "username")
    return None


def _csv_rows(data: bytes) -> list[list[object]]:
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Không thể đọc mã hóa của tệp CSV")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _xlsx_rows(data: bytes) -> list[list[object]]:
    # XLSX is a ZIP container. Bound decompressed size before openpyxl touches it
    # so a small upload cannot expand into an excessive in-memory workbook.
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 2000 or sum(item.file_size for item in infos) > 64 * 1024 * 1024:
                raise ValueError("Tệp XLSX quá lớn sau khi giải nén")
    except zipfile.BadZipFile as exc:
        raise ValueError("Tệp XLSX không hợp lệ") from exc
    wb = load_workbook(
        io.BytesIO(data), read_only=True, data_only=True, keep_links=False
    )
    rows: list[list[object]] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
    finally:
        wb.close()
    return rows


def _xls_rows(data: bytes) -> list[list[object]]:
    book = xlrd.open_workbook(file_contents=data, on_demand=True)
    rows: list[list[object]] = []
    try:
        for sheet in book.sheets():
            for idx in range(sheet.nrows):
                rows.append(sheet.row_values(idx))
    finally:
        book.release_resources()
    return rows


def parse_recipient_file(filename: str, data: bytes, max_recipients: int = 5000) -> dict:
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError("Chỉ hỗ trợ tệp CSV, XLS hoặc XLSX")
    if ext == ".csv":
        rows = _csv_rows(data)
    elif ext == ".xlsx":
        rows = _xlsx_rows(data)
    else:
        rows = _xls_rows(data)

    rows = [row for row in rows if any(str(v or "").strip() for v in row)]
    if not rows:
        raise ValueError("Tệp không có dữ liệu người nhận")

    header_map: dict[int, str] = {}
    for idx, cell in enumerate(rows[0]):
        kind = _header_kind(cell)
        if kind:
            header_map[idx] = kind
    has_header = bool(header_map)
    data_rows = rows[1:] if has_header else rows
    single_col = max((len(r) for r in data_rows), default=0) <= 1

    recipients: list[Recipient] = []
    seen: set[str] = set()
    duplicates = invalid = 0
    for row in data_rows:
        candidates: list[tuple[object, str | None]] = []
        if has_header:
            for col, kind in header_map.items():
                if col < len(row):
                    candidates.append((row[col], kind))
        else:
            candidates.extend((cell, None) for cell in row)
        for cell, hint in candidates:
            if cell is None or not str(cell).strip():
                continue
            rec = normalize_recipient(cell, hint=hint, allow_bare_username=single_col)
            if not rec:
                invalid += 1
                continue
            key = rec.value.lower() if rec.kind == "username" else rec.value
            if key in seen:
                duplicates += 1
                continue
            seen.add(key)
            recipients.append(rec)
            if len(recipients) > max_recipients:
                raise ValueError(f"Danh sách vượt quá giới hạn {max_recipients} người nhận")

    if not recipients:
        raise ValueError("Không tìm thấy số điện thoại quốc tế hoặc username hợp lệ")
    return {
        "recipients": recipients,
        "valid": len(recipients),
        "phones": sum(r.kind == "phone" for r in recipients),
        "usernames": sum(r.kind == "username" for r in recipients),
        "duplicates": duplicates,
        "invalid": invalid,
        "rows": len(data_rows),
    }
